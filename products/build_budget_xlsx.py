# -*- coding: utf-8 -*-
"""오늘살림 가계부 템플릿 (.xlsx) 생성 — 수식 작동, 드롭다운, 대시보드."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.utils import get_column_letter

# ---- 팔레트 (계산해주는 살림 톤: 차분한 세이지+차콜) ----
INK    = "22312B"   # 진한 잉크
SAGE   = "6B8F71"   # 포인트 세이지
SAGE_L = "E8F0EA"   # 연한 세이지 배경
CREAM  = "F7F5EF"   # 크림 배경
LINE   = "D8D8CF"   # 라인
WARN   = "C0563B"   # 초과(빨강)
GOOD   = "4E7A5A"   # 절약(초록)

thin = Side(style="thin", color=LINE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left   = Alignment(horizontal="left", vertical="center")
right  = Alignment(horizontal="right", vertical="center")

def fill(hex_): return PatternFill("solid", fgColor=hex_)
def title_font(sz=11, color=INK, bold=True): return Font(name="맑은 고딕", size=sz, bold=bold, color=color)
def body_font(sz=10, color=INK, bold=False): return Font(name="맑은 고딕", size=sz, bold=bold, color=color)

WON = '#,##0"원"'

CATEGORIES = ["식비", "생활용품", "주거/공과금", "교통", "통신", "의료/건강",
              "문화/여가", "의류/미용", "경조사", "저축/투자", "기타"]
PAY = ["현금", "체크카드", "신용카드", "계좌이체", "간편결제"]
MONTHS = [f"{m}월" for m in range(1, 13)]

wb = openpyxl.Workbook()

# =========================================================
# 1) 표지
# =========================================================
cover = wb.active
cover.title = "표지"
cover.sheet_view.showGridLines = False
for c in "ABCDEFGH":
    cover.column_dimensions[c].width = 12
cover.merge_cells("B2:G3")
t = cover["B2"]; t.value = "오늘살림 가계부"
t.font = Font(name="맑은 고딕", size=28, bold=True, color=INK); t.alignment = center
cover.merge_cells("B4:G4")
s = cover["B4"]; s.value = "계산해주는 살림 · 한 달 -50만원 프로젝트"
s.font = Font(name="맑은 고딕", size=13, color=SAGE); s.alignment = center
for r in range(2, 8):
    for c in range(2, 8):
        cover.cell(row=r, column=c).fill = fill(SAGE_L)

cover.merge_cells("B9:G9")
h = cover["B9"]; h.value = "📌 사용법 (3단계)"
h.font = title_font(13); h.alignment = left
steps = [
    "① [설정] 시트에서 이번 달 카테고리별 예산을 정해요.",
    "② [1월]~[12월] 시트에 그날 쓴 돈을 한 줄씩 적어요. (카테고리는 드롭다운)",
    "③ [대시보드]가 예산 대비 지출·절약액을 자동으로 계산해줘요.",
]
for i, tx in enumerate(steps):
    cover.merge_cells(f"B{10+i}:G{10+i}")
    cell = cover[f"B{10+i}"]; cell.value = tx
    cell.font = body_font(11); cell.alignment = left
cover.merge_cells("B14:G15")
tip = cover["B14"]
tip.value = "💡 완벽하게 다 적으려 하지 마세요. '큰 지출 1개'만 놓치지 않아도 충분해요."
tip.font = Font(name="맑은 고딕", size=11, italic=True, color=SAGE); tip.alignment = left
tip.fill = fill(CREAM)
for r in range(14, 16):
    for c in range(2, 8):
        cover.cell(row=r, column=c).border = border
cover.merge_cells("B17:G17")
c = cover["B17"]; c.value = "ⓒ 오늘살림  ·  개인 사용/1가구용  ·  재판매·재배포 금지"
c.font = body_font(9, "9A9A90"); c.alignment = center

# =========================================================
# 2) 설정 (카테고리·예산)
# =========================================================
st = wb.create_sheet("설정")
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 4
st.column_dimensions["B"].width = 18
st.column_dimensions["C"].width = 16
st.column_dimensions["D"].width = 40
st["B2"] = "카테고리별 월 예산 설정"
st["B2"].font = title_font(14)
hdr = ["카테고리", "월 예산", "메모(선택)"]
for j, h in enumerate(hdr):
    cell = st.cell(row=4, column=2+j, value=h)
    cell.font = title_font(10, "FFFFFF"); cell.fill = fill(SAGE)
    cell.alignment = center; cell.border = border
default_budget = [500000,120000,350000,80000,70000,50000,100000,80000,50000,300000,50000]
for i, (cat, bud) in enumerate(zip(CATEGORIES, default_budget)):
    r = 5+i
    a = st.cell(row=r, column=2, value=cat); a.font = body_font(10); a.border = border; a.alignment = left; a.fill = fill(CREAM)
    b = st.cell(row=r, column=3, value=bud); b.font = body_font(10); b.border = border; b.alignment = right; b.number_format = WON
    d = st.cell(row=r, column=4, value=""); d.border = border
r_tot = 5+len(CATEGORIES)
st.cell(row=r_tot, column=2, value="합계").font = title_font(10)
st.cell(row=r_tot, column=2).fill = fill(SAGE_L); st.cell(row=r_tot, column=2).border = border; st.cell(row=r_tot, column=2).alignment = center
tt = st.cell(row=r_tot, column=3, value=f"=SUM(C5:C{r_tot-1})")
tt.font = title_font(10); tt.fill = fill(SAGE_L); tt.border = border; tt.alignment = right; tt.number_format = WON
st.cell(row=r_tot, column=4).border = border; st.cell(row=r_tot, column=4).fill = fill(SAGE_L)
st["B18"] = "월 목표 저축액"
st["B18"].font = title_font(10)
st["C18"] = 400000; st["C18"].number_format = WON; st["C18"].font = body_font(10); st["C18"].fill = fill(SAGE_L); st["C18"].border = border; st["C18"].alignment = right

# 예산 카테고리 매핑(대시보드 참조용): B5:B15 = 카테고리, C5:C15 = 예산
BUDGET_FIRST, BUDGET_LAST = 5, 5+len(CATEGORIES)-1

# =========================================================
# 3) 월별 입력 시트 (1월~12월)
# =========================================================
INPUT_ROWS = 60  # 한 달 최대 입력 줄
def make_month(name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    widths = [4, 12, 16, 16, 40, 14, 6]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(1+i)].width = w
    ws.merge_cells("B2:E2")
    ws["B2"] = f"{name} 지출 기록"; ws["B2"].font = title_font(14)
    # 요약 박스 (상단 우측)
    ws.merge_cells("F2:G2")
    ws["F2"] = "이번 달 합계"; ws["F2"].font = title_font(10, "FFFFFF"); ws["F2"].fill = fill(SAGE); ws["F2"].alignment = center
    ws.merge_cells("F3:G3")
    ws["F3"] = f"=SUM(F6:F{5+INPUT_ROWS})"; ws["F3"].font = title_font(13, GOOD); ws["F3"].alignment = center; ws["F3"].number_format = WON
    # 헤더
    heads = ["날짜", "카테고리", "결제수단", "내용", "금액", "고정비"]
    for j, h in enumerate(heads):
        cell = ws.cell(row=5, column=2+j, value=h)
        cell.font = title_font(10, "FFFFFF"); cell.fill = fill(INK); cell.alignment = center; cell.border = border
    for r in range(6, 6+INPUT_ROWS):
        for j in range(6):
            cell = ws.cell(row=r, column=2+j); cell.border = border; cell.font = body_font(10)
            if j == 4: cell.number_format = WON; cell.alignment = right
            elif j in (1,2,5): cell.alignment = center
            else: cell.alignment = left
            if r % 2 == 0 and j != -1: cell.fill = fill("FFFFFF")
            else: cell.fill = fill(CREAM)
        ws.cell(row=r, column=2).number_format = "m/d"
    # 드롭다운: 카테고리(C), 결제수단(D), 고정비(G)
    dv_cat = DataValidation(type="list", formula1=f'"{",".join(CATEGORIES)}"', allow_blank=True)
    dv_pay = DataValidation(type="list", formula1=f'"{",".join(PAY)}"', allow_blank=True)
    dv_fix = DataValidation(type="list", formula1='"고정,변동"', allow_blank=True)
    ws.add_data_validation(dv_cat); ws.add_data_validation(dv_pay); ws.add_data_validation(dv_fix)
    dv_cat.add(f"C6:C{5+INPUT_ROWS}")
    dv_pay.add(f"D6:D{5+INPUT_ROWS}")
    dv_fix.add(f"G6:G{5+INPUT_ROWS}")
    return ws

for m in MONTHS:
    make_month(m)

# =========================================================
# 4) 대시보드 (예산 대비 자동 계산)
# =========================================================
db = wb.create_sheet("대시보드", index=2)
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [4,16,16,16,16,14,10]):
    db.column_dimensions[col].width = w
db["B2"] = "대시보드"; db["B2"].font = title_font(16)
db.merge_cells("B2:D2")
# 월 선택
db["F2"] = "조회 월"; db["F2"].font = title_font(10); db["F2"].alignment = right
sel = db["G2"]; sel.value = "1월"; sel.font = title_font(11, SAGE); sel.alignment = center
sel.fill = fill(SAGE_L); sel.border = border
dv_month = DataValidation(type="list", formula1=f'"{",".join(MONTHS)}"', allow_blank=False)
db.add_data_validation(dv_month); dv_month.add("G2")

# 표 헤더
heads = ["카테고리", "예산", "지출", "잔액", "달성률"]
HR = 4
for j, h in enumerate(heads):
    cell = db.cell(row=HR, column=2+j, value=h)
    cell.font = title_font(10, "FFFFFF"); cell.fill = fill(SAGE); cell.alignment = center; cell.border = border
# 각 카테고리 행: 지출 = SUMIF(선택월 시트!카테고리열, 카테고리, 금액열)
for i, cat in enumerate(CATEGORIES):
    r = HR+1+i
    c_cat = db.cell(row=r, column=2, value=cat); c_cat.font = body_font(10); c_cat.border=border; c_cat.alignment=left; c_cat.fill=fill(CREAM)
    # 예산: 설정 시트에서 VLOOKUP
    c_bud = db.cell(row=r, column=3, value=f'=IFERROR(VLOOKUP(B{r},설정!$B${BUDGET_FIRST}:$C${BUDGET_LAST},2,0),0)')
    c_bud.number_format = WON; c_bud.font = body_font(10); c_bud.border=border; c_bud.alignment=right
    # 지출: INDIRECT로 선택월 시트 참조
    c_spent = db.cell(row=r, column=4,
        value=f'=SUMIF(INDIRECT("\'"&$G$2&"\'!C6:C65"),B{r},INDIRECT("\'"&$G$2&"\'!F6:F65"))')
    c_spent.number_format = WON; c_spent.font = body_font(10); c_spent.border=border; c_spent.alignment=right
    # 잔액
    c_left = db.cell(row=r, column=5, value=f"=C{r}-D{r}")
    c_left.number_format = WON; c_left.font = body_font(10); c_left.border=border; c_left.alignment=right
    # 달성률(지출/예산)
    c_rate = db.cell(row=r, column=6, value=f'=IFERROR(D{r}/C{r},0)')
    c_rate.number_format = "0%"; c_rate.font = body_font(10); c_rate.border=border; c_rate.alignment=center
TR = HR+1+len(CATEGORIES)
db.cell(row=TR, column=2, value="합계").font = title_font(10)
for col in range(2,7):
    db.cell(row=TR, column=col).fill = fill(SAGE_L); db.cell(row=TR, column=col).border = border
db.cell(row=TR, column=2).alignment = center
for col, letter in [(3,"C"),(4,"D"),(5,"E")]:
    cc = db.cell(row=TR, column=col, value=f"=SUM({letter}{HR+1}:{letter}{TR-1})")
    cc.number_format = WON; cc.font = title_font(10); cc.alignment = right
rr = db.cell(row=TR, column=6, value=f"=IFERROR(D{TR}/C{TR},0)")
rr.number_format = "0%"; rr.font = title_font(10); rr.alignment = center

# 조건부서식: 달성률 100% 초과 = 빨강, 데이터바
db.conditional_formatting.add(f"F{HR+1}:F{TR-1}",
    CellIsRule(operator="greaterThan", formula=["1"], fill=fill("F4CCC4"), font=Font(color=WARN, bold=True)))
db.conditional_formatting.add(f"F{HR+1}:F{TR-1}",
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=SAGE))
# 잔액 음수 빨강
db.conditional_formatting.add(f"E{HR+1}:E{TR-1}",
    CellIsRule(operator="lessThan", formula=["0"], font=Font(color=WARN, bold=True)))

# KPI 카드 (하단)
K = TR+2
cards = [
    ("이번 달 총지출", f"=D{TR}", GOOD),
    ("총예산 대비 잔액", f"=C{TR}-D{TR}", SAGE),
    ("목표 저축액", "=설정!$C$18", INK),
    ("예산 달성률", f"=IFERROR(D{TR}/C{TR},0)", SAGE),
]
for i,(lab,form,color) in enumerate(cards):
    col = 2 + i
    lc = db.cell(row=K, column=col, value=lab); lc.font = body_font(9,"6A6A62"); lc.alignment=center; lc.fill=fill(SAGE_L); lc.border=border
    vc = db.cell(row=K+1, column=col, value=form); vc.font = title_font(12,color); vc.alignment=center; vc.border=border
    vc.number_format = "0%" if "달성률" in lab else WON
    db.row_dimensions[K+1].height = 26

db["B"+str(K+3)] = "💡 잔액이 빨강이면 그 카테고리는 예산 초과예요. 다음 달 예산을 조정하거나 지출을 점검하세요."
db["B"+str(K+3)].font = Font(name="맑은 고딕", size=10, italic=True, color=SAGE)
db.merge_cells(f"B{K+3}:F{K+3}")

# 시트 순서: 표지, 설정, 대시보드, 1~12월
wb.move_sheet("표지", -(len(wb.sheetnames)))
order = ["표지","설정","대시보드"] + MONTHS
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)

out = "/home/user/memory/products/오늘살림_가계부/오늘살림_가계부_템플릿.xlsx"
wb.save(out)
print("saved:", out)
print("sheets:", wb.sheetnames)
